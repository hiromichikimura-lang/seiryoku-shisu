"""知事の議席台帳(seat registry)。

総務省の年次データを基準点(baseline)として47都道府県ぶん個別に保持し、
政治山で党派欄が確認できた知事選挙の結果だけを都度上書きしていく。
市区町村長・地方議会は総務省側が都道府県単位の集計しか無く、個々の
自治体がどの政党かを特定できないため、同じ仕組みを適用できない
([[seiryoku_shisu_design]]の今後の課題)。
"""
from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import openpyxl

from .fetch import local_parties, seijiyama
from .fetch.util import fetch
from .lineage import canonicalize

REGISTRY_PATH = Path(__file__).resolve().parents[2] / "data" / "governor_registry.json"


def _extract_governor_baseline() -> dict[str, dict]:
    """総務省の年次データから、47都道府県それぞれの現職知事の政党を抽出する。"""
    url, as_of_date = local_parties._find_latest_workbook_url_and_date()
    wb = openpyxl.load_workbook(BytesIO(fetch(url)), data_only=True, read_only=True)
    sheet_name = next(
        n for n in wb.sheetnames if local_parties._classify_sheet(n) == "都道府県知事"
    )
    rows = list(wb[sheet_name].iter_rows(values_only=True))

    header_idx, name_col = next(
        (i, local_parties._find_col(r, "団体名"))
        for i, r in enumerate(rows)
        if r and local_parties._find_col(r, "団体名") is not None
    )
    name_row = rows[header_idx - 1]

    party_cols: dict[str, int] = {}
    for col, val in enumerate(name_row):
        if col < name_col + 2 or not val or val in ("合計", "欠員"):
            continue
        party_cols[val] = col + 2

    baseline = {}
    for row in rows[header_idx + 1 :]:
        if not row or not row[name_col] or row[name_col] in ("合計",):
            continue
        pref = row[name_col]
        party = "無所属"
        for name, col in party_cols.items():
            if col < len(row) and row[col] == 1:
                party = canonicalize(name)
                break
        baseline[pref] = {"party": party, "as_of": as_of_date, "source": "soumu"}
    return baseline


def load_or_init_registry() -> dict[str, dict]:
    if REGISTRY_PATH.exists():
        return json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))
    registry = _extract_governor_baseline()
    save_registry(registry)
    return registry


def save_registry(registry: dict[str, dict]) -> None:
    REGISTRY_PATH.parent.mkdir(exist_ok=True)
    REGISTRY_PATH.write_text(
        json.dumps(registry, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def refresh_from_seijiyama(registry: dict[str, dict], max_pages: int = 5) -> dict[str, dict]:
    """政治山を確認し、党派欄が埋まっている知事選挙の結果だけで台帳を更新する。

    更新するのは registry 内の as_of より新しい投票日の、党派が空欄でない結果のみ。
    """
    updated = dict(registry)
    for page in range(1, max_pages + 1):
        listings = seijiyama.list_elections_page(page)
        if not listings:
            break
        for listing in listings:
            if not listing.name.endswith("知事選挙") or listing.detail_url is None:
                continue
            pref = listing.prefecture
            if pref not in updated:
                continue
            if listing.vote_date <= updated[pref]["as_of"]:
                continue
            candidates = seijiyama.parse_candidates(listing.detail_url)
            winner = next((c for c in candidates if c.elected), None)
            if winner is None or not winner.party:
                continue  # 党派欄が空欄、または未確定の選挙はスキップ
            updated[pref] = {
                "party": canonicalize(winner.party),
                "as_of": listing.vote_date,
                "source": "seijiyama",
            }
    return updated


def current_governor_parties(registry: dict[str, dict]) -> dict[str, int]:
    """{政党名: 都道府県数} の形で、体`都道府県知事`用の議席数辞書を返す。"""
    counts: dict[str, int] = {}
    for entry in registry.values():
        counts[entry["party"]] = counts.get(entry["party"], 0) + 1
    return counts


def coverage_report(registry: dict[str, dict]) -> dict[str, int]:
    """出所別(総務省 vs 政治山で確認済み)の件数を返す。"""
    report: dict[str, int] = {}
    for entry in registry.values():
        report[entry["source"]] = report.get(entry["source"], 0) + 1
    return report
