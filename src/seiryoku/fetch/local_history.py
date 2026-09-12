"""総務省の年次ページを遡って、地方の党派別議席数の時系列を取得する。

一覧ページ(ichiran.html)には年ごとのサブページへのリンクが並んでおり、
新しい年(概ね令和2年以降)はxlsx、それより前(平成22〜令和元年頃)はxls形式で
党派別人員調を配布している。両方をサポートし、それより前の年はPDFや別ページ
構成になり本モジュールでは扱っていない([[seiryoku_shisu_design]]の今後の課題)。
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

import openpyxl
import xlrd
from bs4 import BeautifulSoup

from .local_parties import _classify_sheet, _parse_sheet_totals
from .util import fetch

ICHIRAN_URL = "https://www.soumu.go.jp/senkyo/senkyo_s/data/syozoku/ichiran.html"
_TARGET_LINK_TEXT = "地方公共団体の議会の議員及び長の所属党派別人員調"


@dataclass
class LocalSnapshot:
    label: str  # 一覧ページの表記(例: "令和6年12月31日現在")
    seats: dict[str, dict[str, int]]  # {体: {党派: 議席数}}


def _list_year_pages() -> list[tuple[str, str]]:
    """(表記ラベル, 年次ページURL) のリストを新しい順に返す。"""
    html = fetch(ICHIRAN_URL).decode("shift_jis", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    return [
        (a.get_text(strip=True), "https://www.soumu.go.jp" + a["href"])
        for a in soup.select("dd a")
    ]


def _find_workbook_url(year_page_url: str) -> str | None:
    try:
        html = fetch(year_page_url).decode("shift_jis", errors="ignore")
    except Exception:
        return None
    soup = BeautifulSoup(html, "html.parser")
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if _TARGET_LINK_TEXT in a.get_text(strip=True) and (
            href.endswith(".xlsx") or href.endswith(".xls")
        ):
            return "https://www.soumu.go.jp" + href
    return None


def _sheets_as_rows(url: str) -> dict[str, list[tuple]]:
    """xlsx/xls両対応で、{シート名: 行のリスト} を返す。"""
    content = fetch(url)
    if url.endswith(".xlsx"):
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
        return {name: list(wb[name].iter_rows(values_only=True)) for name in wb.sheetnames}

    wb = xlrd.open_workbook(file_contents=content)
    return {
        sheet.name: [tuple(sheet.row_values(r)) for r in range(sheet.nrows)]
        for sheet in wb.sheets()
    }


def _classify(sheet_name: str, rows: list[tuple]) -> str | None:
    """シート名で分からなければ、シート先頭の見出し文字列で分類する。

    古い年度は「1」「2」...のような通し番号だけがシート名で、
    「（２）都道府県議会議員の所属党派別人員調」のような見出しが
    シートの1行目に入っているため。
    """
    body = _classify_sheet(sheet_name)
    if body is not None:
        return body
    for row in rows[:2]:
        for cell in row or ():
            if cell:
                body = _classify_sheet(str(cell))
                if body is not None:
                    return body
    return None


def _parse_workbook(url: str) -> dict[str, dict[str, int]]:
    merged: dict[str, dict[str, int]] = {}
    for sheet_name, rows in _sheets_as_rows(url).items():
        body = _classify(sheet_name, rows)
        if body is None:
            continue
        totals = _parse_sheet_totals(rows)
        bucket = merged.setdefault(body, {})
        for party, n in totals.items():
            bucket[party] = bucket.get(party, 0) + n
    return merged


def fetch_local_history(max_years: int = 20) -> list[LocalSnapshot]:
    """直近 max_years 年ぶんの地方党派別議席数の時系列を取得する。"""
    snapshots: list[LocalSnapshot] = []
    for label, year_url in _list_year_pages()[:max_years]:
        wb_url = _find_workbook_url(year_url)
        if wb_url is None:
            continue  # PDFのみ等、対応形式が無い年はスキップ
        try:
            seats = _parse_workbook(wb_url)
        except Exception:
            continue
        if seats:
            snapshots.append(LocalSnapshot(label=label, seats=seats))
    return snapshots
