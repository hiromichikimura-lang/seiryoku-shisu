"""総務省「全国地方公共団体コード」改正履歴(市町村合併等の追跡)を取得する。"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

import openpyxl
from bs4 import BeautifulSoup

from .util import fetch

CODE_PAGE_URL = "https://www.soumu.go.jp/denshijiti/code.html"
_TARGET_HEADING_MARKERS = ("改正一覧表", "平成17年4月1日以降")


@dataclass
class CodeChange:
    prefecture: str
    old_code: str | None
    old_name: str | None
    change_type: str | None
    new_code: str | None
    new_name: str | None
    reason: str | None
    effective_date: str | None


def _find_rireki_url() -> str:
    """「改正一覧表(平成17年4月1日以降)」という見出しの直後にあるExcelリンクを探す。

    ダウンロードリンク自体のテキストは全て「Excelファイル」で統一されており
    区別できないため、見出しのテキストを起点に探索する。
    """
    html = fetch(CODE_PAGE_URL).decode("shift_jis", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    heading = soup.find(
        string=lambda s: s and all(marker in s for marker in _TARGET_HEADING_MARKERS)
    )
    if heading is None:
        raise RuntimeError("改正一覧表の見出しが見つかりません")

    node = heading.parent
    for _ in range(20):
        node = node.find_next("a")
        if node is None:
            break
        if node.get("href", "").endswith(".xlsx"):
            return "https://www.soumu.go.jp" + node["href"]
    raise RuntimeError("市区町村コード改正履歴へのリンクが見つかりません")


def fetch_code_changes() -> list[CodeChange]:
    """市区町村コードの改正履歴を一覧として返す(平成17年度以降)。"""
    url = _find_rireki_url()
    wb = openpyxl.load_workbook(BytesIO(fetch(url)), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    changes: list[CodeChange] = []
    current_pref = None
    for row in ws.iter_rows(min_row=5, values_only=True):
        pref, old_code, old_name = row[4], row[5], row[6]
        change_type, new_code, new_name = row[8], row[9], row[10]
        reason, eff_date = row[12], row[13]
        if pref:
            current_pref = pref
        if old_code is None and old_name is None and new_name is None:
            continue
        changes.append(
            CodeChange(
                prefecture=current_pref,
                old_code=str(old_code) if old_code else None,
                old_name=old_name,
                change_type=change_type,
                new_code=str(new_code) if new_code else None,
                new_name=new_name,
                reason=reason,
                effective_date=str(eff_date) if eff_date else None,
            )
        )
    return changes
